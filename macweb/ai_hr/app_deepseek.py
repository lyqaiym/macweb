import argparse
import json
import os

import pdfplumber
from docx import Document
from dotenv import load_dotenv
from openai import OpenAI

load_dotenv()


def read_job_requirement(file_path):
    print("read_job_requirement:path=%s" % file_path)
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def read_resume(file_path):
    print("read_resume:path=%s" % file_path)
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return _read_docx(file_path)
    if ext == ".pdf":
        return _read_pdf(file_path)
    raise ValueError(f"不支持的简历格式：{ext}，仅支持 .docx 和 .pdf")


def _read_docx(file_path):
    document = Document(file_path)
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def _read_pdf(file_path):
    blocks = []
    with pdfplumber.open(file_path) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            blocks.append(f"--- 第 {page_no} 页 ---")
            text = page.extract_text() or ""
            blocks.extend(line.strip() for line in text.splitlines() if line.strip())
            for table in page.extract_tables():
                for row in table:
                    cells = [(c or "").strip() for c in row]
                    if any(cells):
                        blocks.append(" | ".join(cells))
    content = "\n".join(blocks)
    if not content.strip():
        raise ValueError(
            f"{file_path} 未解析出任何文本，可能是扫描版 / 纯图片 PDF，请改用 docx 或可复制文本的 PDF"
        )
    return content


def save_report(file_path, content):
    print("save_report:path=%s" % file_path)
    os.makedirs(os.path.dirname(file_path) or ".", exist_ok=True)
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content)

    return f"评分报告已保存到：{file_path}"


tools = [
    {
        "type": "function",
        "function": {
            "name": "read_job_requirement",
            "description": "读取本地岗位要求（JD）文件内容",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "岗位要求文件路径，txt / md 等纯文本"
                    }
                },
                "required": ["file_path"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "read_resume",
            "description": "读取候选人简历（支持 docx 和 pdf），返回其中的正文段落和表格文本",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "简历文件路径，扩展名为 .docx 或 .pdf"
                    }
                },
                "required": ["file_path"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "save_report",
            "description": "把生成的简历评分报告保存到本地文件",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "保存文件路径"
                    },
                    "content": {
                        "type": "string",
                        "description": "完整的评分报告内容，markdown 格式"
                    }
                },
                "required": ["file_path", "content"],
                "additionalProperties": False
            }
        }
    }
]

TOOL_CALL_MAP = {
    "read_job_requirement": read_job_requirement,
    "read_resume": read_resume,
    "save_report": save_report,
}

SYSTEM_PROMPT = """
你是一名资深技术招聘专家（HR + 技术面评委）。
你的任务是：读取岗位要求和候选人简历，严格按岗位要求给简历打分。

评分维度与权重（总分 100）：
- 硬性条件（学历、工作年限、必备技能、地点/语言等硬门槛）：25 分
- 技术栈匹配度（岗位要求的技术与简历实际使用深度）：25 分
- 项目 / 业务经验匹配度（规模、复杂度、行业相关性、个人贡献）：25 分
- 职业发展与稳定性（跳槽频率、职级成长、空窗期）：15 分
- 加分项（开源、专利、论文、竞赛、影响力、岗位额外偏好）：10 分

打分要求：
- 每个维度给出得分、满分、评分依据，依据必须引用简历中的原文或具体事实，禁止臆造经历。
- 简历中没有体现的能力按"未体现"处理并扣分，不要假设候选人具备。
- 若命中硬性条件不达标（如学历、年限、必备技能缺失），在报告开头明确标注"硬性条件不满足项"。
- 总分为各维度得分之和；根据总分给出结论：
  85 以上强烈推荐面试 / 70-84 推荐面试 / 60-69 待定 / 60 以下不推荐。

报告结构（markdown）：
1. 候选人概览（姓名或脱敏标识、当前岗位、总年限、学历）
2. 总分与结论
3. 各维度评分表（维度 / 得分 / 满分 / 评分依据）
4. 匹配亮点
5. 风险与差距
6. 建议面试追问问题（3-5 条，针对简历中含糊或高风险的点）
""".strip()


def run_turn(client, model, turn, messages):
    sub_turn = 1
    while True:
        response = client.chat.completions.create(
            model=model,
            messages=messages,
            tools=tools,
            reasoning_effort="high",
            extra_body={"thinking": {"type": "enabled"}},
        )
        msg = response.choices[0].message
        # reasoning_content 不能回传给下一轮，否则接口报 400
        assistant_message = msg.model_dump(exclude_none=True)
        assistant_message.pop("reasoning_content", None)
        messages.append(assistant_message)
        reasoning_content = getattr(msg, "reasoning_content", None)
        content = msg.content
        tool_calls = msg.tool_calls
        print(f"Turn {turn}.{sub_turn}\n{reasoning_content=}\n{content=}\n{tool_calls=}")
        if not tool_calls:
            break
        for tool in tool_calls:
            tool_function = TOOL_CALL_MAP[tool.function.name]
            tool_result = tool_function(**json.loads(tool.function.arguments))
            print(f"tool result for {tool.function.name}: {tool_result}\n")
            messages.append({
                "role": "tool",
                "tool_call_id": tool.id,
                "content": tool_result,
            })
        sub_turn += 1
    print()


def main():
    parser = argparse.ArgumentParser(description="简历打分系统：按岗位要求给简历打分")
    parser.add_argument("--jd", required=True, help="岗位要求文件路径（txt / md）")
    parser.add_argument("--resume", required=True, help="候选人简历路径（docx / pdf）")
    parser.add_argument("--out", help="评分报告输出路径，默认 output/<简历名>_score.md")
    parser.add_argument("--model", default="deepseek-v4-pro", help="模型名")
    args = parser.parse_args()

    out_path = args.out or os.path.join(
        "output", os.path.splitext(os.path.basename(args.resume))[0] + "_score.md"
    )

    client = OpenAI(
        api_key=os.environ.get("DEEPSEEK_API_KEY"),
        base_url=os.environ.get("DEEPSEEK_BASE_URL"),
    )

    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {
            "role": "user",
            "content": (
                f"岗位要求文件：{args.jd}\n"
                f"候选人简历：{args.resume}\n"
                f"请读取这两个文件，完成打分，并把 markdown 报告保存到：{out_path}"
            ),
        },
    ]
    run_turn(client, args.model, 1, messages)


if __name__ == "__main__":
    main()
